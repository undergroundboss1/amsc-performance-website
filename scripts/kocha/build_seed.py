"""
Build scripts/supabase-kocha-seed-v1.sql from the Kocha Mwenzangu sources.

    python3 scripts/kocha/build_seed.py

Inputs:
  scripts/kocha/content.json                       weeks, study cards, homework,
                                                   files, quick cards, glossary,
                                                   WhatsApp templates
  docs/kocha-mwenzangu-source/
    AMSC_Kocha_Mwenzangu_Video_Shortlist.xlsx      every shortlisted video

The generated SQL only INSERTs rows that are missing (ON CONFLICT DO
NOTHING), so re-running it never overwrites what Arnold has edited or
approved in the admin portal. Every video is imported as 'pending': only
Arnold's approval makes one visible to scholars.

Needs openpyxl (already in requirements.txt).
"""

import json
import re
from datetime import date, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
CONTENT = ROOT / 'scripts' / 'kocha' / 'content.json'
SHORTLIST = ROOT / 'docs' / 'kocha-mwenzangu-source' / 'AMSC_Kocha_Mwenzangu_Video_Shortlist.xlsx'
OUT = ROOT / 'scripts' / 'supabase-kocha-seed-v1.sql'

# Shortlist "Topic / movement pattern" wording -> kocha.videos.movement_pattern.
PATTERNS = {
    'squat': 'squat',
    'hip hinge': 'hinge',
    'hinge': 'hinge',
    'lunge and split squat': 'lunge',
    'single-leg strength and balance': 'single-leg',
    'push': 'push',
    'pull (row)': 'pull',
    'core stability': 'core',
    'carry': 'carry',
    'jumping and landing': 'jumping-landing',
    'plyometrics': 'jumping-landing',
    'stopping and changing direction': 'stopping-cod',
    'mobility': 'mobility',
    'warm-up': 'warm-up',
    'testing': 'testing',
    'conditioning': 'energy-systems',
    'nutrition': 'nutrition',
    'hydration': 'hydration',
    'return to play': 'return-to-play',
}

# Rows with no topic take the pattern their session is about.
SESSION_PATTERN = {2: 'testing', 5: 'energy-systems', 8: 'sleep', 10: 'return-to-play'}

NUTRITION_PATTERNS = {'nutrition', 'hydration'}
NUTRITION_TOPICS = {'game day'}  # tournament-day eating
NUTRITION_WEEKS = {7, 8}


def q(value):
    """SQL literal."""
    if value is None:
        return 'NULL'
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, int):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def q_text_array(values):
    return 'ARRAY[' + ', '.join(q(v) for v in values) + ']::text[]' if values else "'{}'::text[]"


def q_int_array(values):
    return 'ARRAY[' + ', '.join(str(int(v)) for v in values) + ']::smallint[]' if values else "'{}'::smallint[]"


def q_json(value):
    return 'NULL' if value is None else q(json.dumps(value, ensure_ascii=False)) + '::jsonb'


def eat(day, hhmm='00:00'):
    """A wall-clock time in Nairobi (EAT, UTC+3, no daylight saving)."""
    return q(f'{day.isoformat()} {hhmm}:00+03')


def cohort_id(slug):
    return f"(SELECT id FROM kocha.cohorts WHERE slug = {q(slug)})"


def youtube_id(url):
    m = re.search(r'(?:v=|youtu\.be/)([A-Za-z0-9_-]{11})', url or '')
    return m.group(1) if m else None


def week_rows(content):
    slug = content['cohort']['slug']
    start = date.fromisoformat(content['cohort']['starts_on'])
    rel = content['release']
    out = []
    for w in content['weeks']:
        class_on = start + timedelta(weeks=w['n'] - 1)
        release_on = class_on - timedelta(days=rel['open_days_before_class'])
        due_on = class_on + timedelta(days=rel['due_days_after_class'])
        gate = w.get('gate')
        card = w.get('study_card')
        if card is None and w.get('summary_focus'):
            # Weeks 9 to 12: the full card is issued as the session nears.
            card = {'focus': w['summary_focus']}
        out.append(
            'INSERT INTO kocha.weeks (cohort_id, n, title, takeaway, competencies, class_on, joint, '
            'release_at, due_at, published, gate, study_card)\n'
            f"SELECT id, {w['n']}, {q(w['title'])}, {q(w.get('takeaway'))}, {q_text_array(w['competencies'])}, "
            f"{q(class_on.isoformat())}, {q(bool(w.get('joint')))}, {eat(release_on)}, {eat(due_on, rel['due_time'])}, "
            # Gated weeks are double-locked: unpublished AND behind the gate.
            f"{q(gate is None)}, {q(gate)}, {q_json(card)}\n"
            f"FROM kocha.cohorts WHERE slug = {q(slug)}\n"
            'ON CONFLICT (cohort_id, n) DO NOTHING;'
        )
    return out


def assignment_rows(content):
    slug = content['cohort']['slug']
    out = []
    for a in content['assignments']:
        out.append(
            'INSERT INTO kocha.assignments (week_id, title, time_needed, brief, submit_what, done_checklist, example, optional)\n'
            f"SELECT w.id, {q(a['title'])}, {q(a.get('time_needed'))}, {q(a['brief'])}, {q(a['submit_what'])}, "
            f"{q_text_array(a.get('done_checklist', []))}, {q(a.get('example'))}, {q(bool(a.get('optional')))}\n"
            'FROM kocha.weeks w JOIN kocha.cohorts c ON c.id = w.cohort_id\n'
            f"WHERE c.slug = {q(slug)} AND w.n = {a['week']}\n"
            'ON CONFLICT (week_id) DO NOTHING;'
        )
    return out


def file_rows(content):
    slug = content['cohort']['slug']
    out = []
    for f in content['files']:
        # Seeded unpublished with no storage path: the object is uploaded from
        # the admin Files tab, which sets the path and size, then Arnold
        # publishes it.
        out.append(
            'INSERT INTO kocha.files (cohort_id, slug, title, description, kind, format, size_bytes, weeks, '
            'show_on_home, gate, published, sort)\n'
            f"VALUES ({cohort_id(slug)}, {q(f['slug'])}, {q(f['title'])}, {q(f.get('description'))}, {q(f['kind'])}, "
            f"{q(f['format'])}, {q(f.get('size_bytes'))}, {q_int_array(f.get('weeks', []))}, "
            f"{q(bool(f.get('show_on_home')))}, {q(f.get('gate'))}, FALSE, {q(f.get('sort', 0))})\n"
            'ON CONFLICT (cohort_id, slug) DO NOTHING;'
        )
    return out


def simple_rows(content):
    slug = content['cohort']['slug']
    cid = cohort_id(slug)
    out = []
    for i, c in enumerate(content['quick_cards']):
        out.append(
            'INSERT INTO kocha.quick_cards (cohort_id, slug, title, body, gate, sort)\n'
            f"VALUES ({cid}, {q(c['slug'])}, {q(c['title'])}, {q_json(c['body'])}, {q(c.get('gate'))}, {(i + 1) * 10})\n"
            'ON CONFLICT (cohort_id, slug) DO NOTHING;'
        )
    for g in content['glossary']:
        out.append(
            'INSERT INTO kocha.glossary (cohort_id, term, meaning, first_week, gate)\n'
            f"VALUES ({cid}, {q(g['term'])}, {q(g['meaning'])}, {q(g.get('first_week'))}, {q(g.get('gate'))})\n"
            'ON CONFLICT (cohort_id, term) DO NOTHING;'
        )
    for i, t in enumerate(content['message_templates']):
        out.append(
            'INSERT INTO kocha.message_templates (cohort_id, key, title, body, sort)\n'
            f"VALUES ({cid}, {q(t['key'])}, {q(t['title'])}, {q(t['body'])}, {(i + 1) * 10})\n"
            'ON CONFLICT (cohort_id, key) DO NOTHING;'
        )
    return out


def video_rows():
    ws = openpyxl.load_workbook(SHORTLIST, data_only=True).active
    header = [str(c or '').strip().lower() for c in next(ws.iter_rows(max_row=1, values_only=True))]
    col = {name: header.index(name) for name in header if name}
    out = []
    gaps = 0
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if not any(row):
            continue
        get = lambda name: (str(row[col[name]]).strip() if row[col[name]] is not None else None)
        session = get('session')
        weeks = [] if session in (None, 'all') else [int(float(session))]
        topic = get('topic / movement pattern')
        status_text = get('status') or ''
        title = get('title')
        why = get('why / note')
        length = get('length')
        length = None if length in (None, 'n/a') else length

        pattern = PATTERNS.get((topic or '').lower())
        if pattern is None and topic is None and weeks:
            n = weeks[0]
            if n == 7:
                pattern = 'hydration' if 'hydration' in (title or '').lower() else 'nutrition'
            else:
                pattern = SESSION_PATTERN.get(n)

        gate = 'nutrition' if (
            set(weeks) & NUTRITION_WEEKS
            or pattern in NUTRITION_PATTERNS
            or (topic or '').lower() in NUTRITION_TOPICS
        ) else None

        deeper = 'deep dive' in (why or '').lower()

        if status_text.upper().startswith('GAP'):
            gaps += 1
            # Reserved slot for an AMSC-made video: a placeholder Arnold can
            # fill later (link + title), then approve.
            out.append(
                'INSERT INTO kocha.videos (source_key, kind, title, weeks, topic, movement_pattern, status, '
                'shortlist_status, gate, notes, added_by, sort)\n'
                f"VALUES ({q(f'gap:{gaps}')}, 'amsc', {q(f'AMSC video: {topic}' if topic else 'AMSC video')}, "
                f"{q_int_array(weeks)}, {q(topic)}, {q(pattern)}, 'pending', {q(status_text)}, {q(gate)}, "
                f"{q('No suitable outside video found. Reserved for an AMSC-made video.')}, 'shortlist', {i * 10})\n"
                'ON CONFLICT (source_key) DO NOTHING;'
            )
            continue

        url = get('url')
        yid = youtube_id(url)
        if not yid:
            raise SystemExit(f'Shortlist row {i + 1}: no YouTube id in {url!r}')
        out.append(
            'INSERT INTO kocha.videos (source_key, kind, youtube_id, url, title, channel, length_text, why_watch, '
            'weeks, topic, movement_pattern, audience_level, status, shortlist_status, gate, notes, added_by, sort)\n'
            f"VALUES ({q('yt:' + yid)}, 'youtube', {q(yid)}, {q(url)}, {q(title)}, {q(get('channel'))}, {q(length)}, "
            f"{q(why)}, {q_int_array(weeks)}, {q(topic)}, {q(pattern)}, {q('deeper' if deeper else 'beginner')}, "
            f"'pending', {q(status_text)}, {q(gate)}, {q(get('arnold notes'))}, 'shortlist', {i * 10})\n"
            'ON CONFLICT (source_key) DO NOTHING;'
        )
    return out


def main():
    content = json.loads(CONTENT.read_text(encoding='utf-8'))
    c = content['cohort']
    sections = [
        ('Cohort', [
            'INSERT INTO kocha.cohorts (slug, name, starts_on)\n'
            f"VALUES ({q(c['slug'])}, {q(c['name'])}, {q(c['starts_on'])})\n"
            'ON CONFLICT (slug) DO NOTHING;'
        ]),
        ('12-week map and study cards', week_rows(content)),
        ('Homework assignments', assignment_rows(content)),
        ('File manifest (uploaded and published from the admin Files tab)', file_rows(content)),
        ('Quick-reference cards, glossary, WhatsApp templates', simple_rows(content)),
        ('Videos from the shortlist (all pending until approved)', video_rows()),
    ]
    lines = [
        '-- =============================================================',
        '-- AMSC Performance — Kocha Mwenzangu seed v1 (GENERATED — do not edit)',
        '-- =============================================================',
        '-- Generated by scripts/kocha/build_seed.py from scripts/kocha/content.json',
        '-- and docs/kocha-mwenzangu-source/AMSC_Kocha_Mwenzangu_Video_Shortlist.xlsx.',
        '-- Edit those sources and regenerate instead of editing this file.',
        '--',
        '-- Run AFTER scripts/supabase-kocha-schema-v1.sql.',
        '--',
        '-- Safe to re-run: every INSERT is ON CONFLICT DO NOTHING, so rows that',
        '-- already exist — including anything edited or approved in the admin',
        '-- portal — are left exactly as they are.',
        '-- =============================================================',
        '',
        'BEGIN;',
    ]
    for title, rows in sections:
        lines += ['', f'-- ---------- {title} ----------', '']
        lines += [r + '\n' for r in rows]
    lines += ['COMMIT;', '']
    OUT.write_text('\n'.join(lines), encoding='utf-8')
    print(f'Wrote {OUT.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
